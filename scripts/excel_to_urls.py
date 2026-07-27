"""
excel_to_urls.py - Lit un ou plusieurs fichiers Excel produits par wiki_crawler.py
et ajoute les URLs marquees "O" dans urls.txt (sans doublons).

Usage :
    python excel_to_urls.py <fichier.xlsx> [fichier2.xlsx ...]

    Par defaut lit tous les .xlsx dans output/ et ecrit dans urls.txt
"""

import sys
from pathlib import Path
import openpyxl

OUTPUT_DIR = Path("output")
DEFAULT_URLS_FILE = Path("urls.txt")


def is_fully_filled(xlsx_path: Path) -> bool:
    """Retourne True si toutes les lignes ont une valeur O/N dans Inclure."""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active
    for row in ws.iter_rows(min_row=2, values_only=True):
        if len(row) < 4:
            continue
        inclure = str(row[3]).strip().upper() if row[3] is not None else ""
        if inclure not in ("O", "N"):
            wb.close()
            return False
    wb.close()
    return True


def extract_urls(xlsx_path: Path) -> list[str]:
    """Retourne les URLs marquees O (pages a traiter au prochain tour)."""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active
    urls = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if len(row) < 4:
            continue
        url = row[1]
        inclure = str(row[3]).strip().upper() if row[3] is not None else ""
        if url and inclure == "O":
            urls.append(str(url).strip())
    wb.close()
    return urls


def load_existing_urls(urls_file: Path) -> set[str]:
    """Retourne toutes les URLs deja presentes dans urls.txt (commentees ou non)."""
    if not urls_file.exists():
        return set()
    lines = urls_file.read_text(encoding="utf-8").splitlines()
    result = set()
    for line in lines:
        stripped = line.strip().lstrip("# ").strip()
        if stripped:
            result.add(stripped)
    return result


def main() -> None:
    args = sys.argv[1:]

    # Fichiers Excel a lire
    if args:
        xlsx_files = [Path(a) for a in args]
    else:
        xlsx_files = sorted(OUTPUT_DIR.glob("*.xlsx"))

    if not xlsx_files:
        print("[X] Aucun fichier Excel trouve.")
        sys.exit(1)

    urls_file = DEFAULT_URLS_FILE
    existing = load_existing_urls(urls_file)

    new_urls: list[str] = []
    seen: set[str] = set(existing)

    for xlsx in xlsx_files:
        if not xlsx.exists():
            print(f"  [!] Fichier introuvable : {xlsx}")
            continue
        if not is_fully_filled(xlsx):
            print(f"  [~] {xlsx.name} : lignes sans decision, ignore")
            continue
        urls = extract_urls(xlsx)
        added = []
        for u in urls:
            if u not in seen:
                seen.add(u)
                new_urls.append(u)
                added.append(u)
        print(f"  [OK] {xlsx.name} -> {len(urls)} O, {len(added)} nouvelles")

    if not new_urls:
        print("[!] Aucune nouvelle URL a ajouter.")
        sys.exit(0)

    with urls_file.open("a", encoding="utf-8") as f:
        f.write("\n")
        for url in new_urls:
            f.write(url + "\n")

    print(f"\n[OK] {len(new_urls)} URL(s) ajoutees dans {urls_file}")


if __name__ == "__main__":
    main()
