"""
dedup_xlsx.py - Deduplique les URLs entre tous les fichiers Excel de output/

Pour chaque URL qui apparait dans plusieurs fichiers, seule la premiere occurrence
(fichier traite en ordre alphabetique) est conservee. Les doublons sont supprimes.

Usage :
    python dedup_xlsx.py [dossier]

    Par defaut traite tous les .xlsx dans output/
"""

import sys
from pathlib import Path
import openpyxl

OUTPUT_DIR = Path("output")


def dedup(folder: Path) -> None:
    xlsx_files = sorted(folder.glob("*.xlsx"))

    if not xlsx_files:
        print("[X] Aucun fichier Excel trouve.")
        sys.exit(1)

    print(f"[OK] {len(xlsx_files)} fichier(s) a analyser\n")

    seen_urls: set[str] = set()
    total_removed = 0

    for xlsx in xlsx_files:
        wb = openpyxl.load_workbook(xlsx)
        ws = wb.active

        rows_to_delete = []
        for row_idx in range(2, ws.max_row + 1):
            url = ws.cell(row=row_idx, column=2).value
            if not url:
                continue
            url = str(url).strip()
            if url in seen_urls:
                rows_to_delete.append(row_idx)
            else:
                seen_urls.add(url)

        if rows_to_delete:
            # Supprimer en ordre inverse pour ne pas decaler les indices
            for row_idx in reversed(rows_to_delete):
                ws.delete_rows(row_idx)
            wb.save(xlsx)
            total_removed += len(rows_to_delete)
            print(f"  {xlsx.name} : {len(rows_to_delete)} doublon(s) supprime(s)")
        else:
            print(f"  {xlsx.name} : aucun doublon")

    print(f"\n[OK] {total_removed} doublon(s) supprime(s) au total")


def main() -> None:
    folder = Path(sys.argv[1]) if len(sys.argv) > 1 else OUTPUT_DIR
    if not folder.is_dir():
        print(f"[X] Dossier introuvable : {folder}")
        sys.exit(1)
    dedup(folder)


if __name__ == "__main__":
    main()
