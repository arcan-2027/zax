"""
wiki_crawler.py - Step 1 : Decouverte de liens internes par section

Usage :
    python wiki_crawler.py [fichier_urls.txt]

    Par defaut lit "urls.txt" dans le repertoire courant.

Input :
    Fichier texte, une URL de page fallout-wiki.com par ligne.

Output :
    Un fichier Excel par page dans le dossier output/
    Colonnes : Nom de la page | URL | Section | Inclure

Deduplication : premiere occurrence uniquement.
Filtres : redlinks, namespaces hors-lore (Special, Categorie, Aide, etc.),
          liens d'action (edit, history...), liens externes.
"""

import sys
import time
import re
from pathlib import Path

import requests
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from dedup_xlsx import dedup

# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------

BASE_URL = "https://fallout-wiki.com"
OUTPUT_DIR = Path("output")
DEFAULT_INPUT = Path("urls.txt")
REQUEST_DELAY = 0.75  # secondes entre chaque requete (politesse serveur)

# Prefixes de chemins a ignorer (namespaces hors-lore)
FILTERED_PREFIXES = (
    "Sp%C3%A9cial:",  # Special: encode
    "Sp%C3%A3cial:",
    "Spécial:",
    "Catégorie:",
    "Cat%C3%A9gorie:",
    "Aide:",
    "Fichier:",
    "File:",
    "Discussion:",
    "Portail:",
    "Portail%3A",
    "Mod%C3%A8le:",
    "Modèle:",
    "Les_Archives_de_Vault-Tec:",
    "Les%20Archives",
    "Utilisateur:",
    "Utilisateur_discussion:",
    "MediaWiki:",
    "Special:",
    "Category:",
    "Help:",
    "Template:",
    "Talk:",
)

# Styles Excel (palette ZAX)
HEADER_BG = "0A5C0A"
HEADER_FG = "39FF14"
ROW_BG_ALT = "071407"   # rangees alternees
ROW_BG_NRM = "0A0A0A"
ROW_FG = "1DB308"
BORDER_COLOR = "1A5C1A"

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def sanitize_filename(name: str) -> str:
    """Retire les caracteres interdits dans un nom de fichier Windows."""
    name = re.sub(r'[<>:"/\\|?*]', '_', name)
    name = name.strip(". ")
    return name or "page_sans_titre"


def is_filtered(href: str) -> bool:
    """Retourne True si le lien doit etre ignore."""
    path = href.lstrip("/")
    return any(path.startswith(prefix) for prefix in FILTERED_PREFIXES)


def extract_section_name(h_tag) -> str:
    """Extrait le texte d'un tag h2/h3 en ignorant le lien [modifier]."""
    span = h_tag.find("span", class_="mw-headline")
    if span:
        return span.get_text(strip=True)
    # Fallback : texte brut sans le sous-tag [modifier]
    for child in h_tag.find_all("span", class_="mw-editsection"):
        child.decompose()
    return h_tag.get_text(strip=True)

# ---------------------------------------------------------------------------
# Scraping
# ---------------------------------------------------------------------------

def fetch_page(url: str) -> BeautifulSoup:
    headers = {"User-Agent": "ZAX-wiki-crawler/1.0 (GN Fallout 2027 lore tool)"}
    resp = requests.get(url, headers=headers, timeout=20)
    resp.raise_for_status()
    return BeautifulSoup(resp.text, "html.parser")


def get_links_by_section(soup: BeautifulSoup) -> tuple[str, list[dict]]:
    """
    Retourne (titre_h1, liste_de_liens).
    Chaque lien : {"name": str, "url": str, "section": str}
    Deduplication par URL (premiere occurrence).
    """
    # Titre de la page
    h1_tag = soup.find("h1", id="firstHeading") or soup.find("h1")
    page_title = h1_tag.get_text(strip=True) if h1_tag else "Inconnu"

    # Zone de contenu principale
    content = soup.find(id="mw-content-text")
    if not content:
        return page_title, []

    # Suppression des elements parasites avant parsing
    for element in content.find_all(id="toc"):          # table des matieres
        element.decompose()
    for element in content.find_all(class_="navbox"):   # boites de navigation
        element.decompose()
    for element in content.find_all(class_="mw-editsection"):  # liens [modifier]
        element.decompose()
    for element in content.find_all(class_="reference"):       # notes de bas de page
        element.decompose()
    for element in content.find_all("sup"):             # exposants (refs)
        element.decompose()

    links = []
    seen_urls: set[str] = set()
    current_section = "(intro)"

    for element in content.descendants:
        # NavigableString : pas de .name utilisable
        if not hasattr(element, "name") or element.name is None:
            continue

        # Mise a jour de la section courante
        if element.name in ("h2", "h3"):
            current_section = extract_section_name(element)
            continue

        # Traitement des liens
        if element.name != "a":
            continue

        href = element.get("href", "")
        if not href:
            continue

        # Rejets rapides
        if "redlink=1" in href:
            continue
        if not href.startswith("/") or href.startswith("//"):
            continue
        if "action=" in href:
            continue
        if "index.php" in href and "action" not in href:
            # Liens de recherche ou autres pages speciales
            if "title=" in href:
                continue
        if is_filtered(href):
            continue

        # Deduplication (sans fragment d'ancre)
        clean_href = href.split("#")[0]
        if not clean_href or clean_href in seen_urls:
            continue
        seen_urls.add(clean_href)

        link_name = element.get_text(strip=True)
        if not link_name:
            continue

        links.append({
            "name": link_name,
            "url": BASE_URL + href,
            "section": current_section,
        })

    return page_title, links

# ---------------------------------------------------------------------------
# Export Excel
# ---------------------------------------------------------------------------

def thin_border() -> Border:
    side = Side(style="thin", color=BORDER_COLOR)
    return Border(left=side, right=side, top=side, bottom=side)


def write_excel(page_title: str, links: list[dict]) -> Path:
    OUTPUT_DIR.mkdir(exist_ok=True)
    filename = sanitize_filename(page_title) + ".xlsx"
    filepath = OUTPUT_DIR / filename

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Liens"

    # En-tetes
    headers = ["Nom de la page", "URL", "Section", "Inclure"]
    header_font = Font(bold=True, color=HEADER_FG, name="Courier New")
    header_fill = PatternFill("solid", fgColor=HEADER_BG)
    header_align = Alignment(horizontal="center", vertical="center")

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border()

    ws.row_dimensions[1].height = 20

    # Donnees
    normal_align = Alignment(vertical="center", wrap_text=False)
    url_font = Font(color="0A5C0A", name="Courier New", size=9)

    for row_idx, link in enumerate(links, 2):
        bg = ROW_BG_ALT if row_idx % 2 == 0 else ROW_BG_NRM
        row_fill = PatternFill("solid", fgColor=bg)
        row_font = Font(color=ROW_FG, name="Courier New", size=10)

        ws.cell(row=row_idx, column=1, value=link["name"]).font = row_font
        ws.cell(row=row_idx, column=1).fill = row_fill
        ws.cell(row=row_idx, column=1).alignment = normal_align
        ws.cell(row=row_idx, column=1).border = thin_border()

        url_cell = ws.cell(row=row_idx, column=2, value=link["url"])
        url_cell.font = url_font
        url_cell.fill = row_fill
        url_cell.alignment = normal_align
        url_cell.border = thin_border()

        ws.cell(row=row_idx, column=3, value=link["section"]).font = row_font
        ws.cell(row=row_idx, column=3).fill = row_fill
        ws.cell(row=row_idx, column=3).alignment = normal_align
        ws.cell(row=row_idx, column=3).border = thin_border()

        # Colonne Inclure : validation O/N
        inclure_cell = ws.cell(row=row_idx, column=4, value="")
        inclure_cell.font = Font(color=HEADER_FG, bold=True, name="Courier New", size=11)
        inclure_cell.fill = PatternFill("solid", fgColor="071407")
        inclure_cell.alignment = Alignment(horizontal="center", vertical="center")
        inclure_cell.border = thin_border()

    # Largeurs de colonnes
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 62
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["D"].width = 10

    # Figer la ligne d'en-tete
    ws.freeze_panes = "A2"

    # Hauteur des lignes de donnees
    for row_idx in range(2, len(links) + 2):
        ws.row_dimensions[row_idx].height = 16

    wb.save(filepath)
    return filepath

# ---------------------------------------------------------------------------
# Point d'entree
# ---------------------------------------------------------------------------

def comment_url_in_file(input_file: Path, url: str) -> None:
    """Remplace la ligne contenant l'URL par # <url> dans le fichier source."""
    lines = input_file.read_text(encoding="utf-8").splitlines()
    new_lines = []
    for line in lines:
        if line.strip() == url:
            new_lines.append(f"# {line.strip()}")
        else:
            new_lines.append(line)
    input_file.write_text("\n".join(new_lines) + "\n", encoding="utf-8")


def main() -> None:
    input_file = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_INPUT

    if not input_file.exists():
        print(f"[X] Fichier introuvable : {input_file}")
        print(f"    Usage : python wiki_crawler.py [fichier_urls.txt]")
        sys.exit(1)

    raw_lines = input_file.read_text(encoding="utf-8").splitlines()
    urls = [line.strip() for line in raw_lines if line.strip() and not line.startswith("#")]

    if not urls:
        print("[X] Aucune URL trouvee dans le fichier.")
        sys.exit(1)

    print(f"[OK] {len(urls)} URL(s) a traiter -> dossier output/")
    print()

    for i, url in enumerate(urls, 1):
        print(f"[{i}/{len(urls)}] {url}")
        try:
            soup = fetch_page(url)
            title, links = get_links_by_section(soup)

            if not links:
                print(f"  [!] Aucun lien interne trouve sur cette page.")
            else:
                filepath = write_excel(title, links)
                print(f"  [OK] \"{title}\" -> {len(links)} liens -> {filepath.name}")

            comment_url_in_file(input_file, url)

        except requests.HTTPError as e:
            print(f"  [X] Erreur HTTP {e.response.status_code} : {url}")
        except requests.RequestException as e:
            print(f"  [X] Erreur reseau : {e}")
        except Exception as e:
            print(f"  [X] Erreur inattendue : {e}")

        if i < len(urls):
            time.sleep(REQUEST_DELAY)

    print()
    print("[OK] Deduplication des xlsx...")
    dedup(OUTPUT_DIR)
    print("[OK] Termine.")


if __name__ == "__main__":
    main()
